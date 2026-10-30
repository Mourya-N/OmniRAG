"""PowerPoint (.pptx) Presentation Parser using python-pptx."""

import io
import uuid
from typing import Any

from pptx import Presentation

from app.core.exceptions import ParserError
from app.services.parser.base import DocumentParser
from app.services.parser.schema import (
    CanonicalDocument,
    ContentBlock,
    PageData,
    TableData,
)


class PPTXParser(DocumentParser):
    """Parses PowerPoint presentations (.pptx) into CanonicalDocument."""

    def can_parse(self, mime_type: str) -> bool:
        return mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    def parse(self, content: bytes, filename: str, document_id: str = "") -> CanonicalDocument:
        try:
            prs = Presentation(io.BytesIO(content))
        except Exception as e:
            raise ParserError(f"Failed to open PPTX stream: {e}") from e

        doc_id = document_id or str(uuid.uuid4())
        pages: list[PageData] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            page_data = PageData(page_number=slide_num)

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            page_data.blocks.append(
                                ContentBlock(
                                    block_id=str(uuid.uuid4()),
                                    block_type="text",
                                    text=text,
                                )
                            )

                if shape.has_table:
                    tbl = shape.table
                    headers = [cell.text.strip() for cell in tbl.rows[0].cells]
                    rows: list[list[str]] = []
                    for row in tbl.rows[1:]:
                        rows.append([cell.text.strip() for cell in row.cells])

                    page_data.tables.append(
                        TableData(
                            table_id=str(uuid.uuid4()),
                            headers=headers,
                            rows=rows,
                            page_number=slide_num,
                        )
                    )

            pages.append(page_data)

        return CanonicalDocument(
            document_id=doc_id,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            total_pages=len(pages),
            pages=pages,
            metadata={"format": "PPTX", "slides_count": len(pages)},
        )
