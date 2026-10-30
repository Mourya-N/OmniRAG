"""DOCX Document Parser using python-docx."""

import io
import uuid
from typing import Any

import docx

from app.core.exceptions import ParserError
from app.services.parser.base import DocumentParser
from app.services.parser.schema import (
    CanonicalDocument,
    ContentBlock,
    PageData,
    TableData,
)


class DocxParser(DocumentParser):
    """Parses Microsoft Word (.docx) files into CanonicalDocument structures."""

    def can_parse(self, mime_type: str) -> bool:
        return mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def parse(self, content: bytes, filename: str, document_id: str = "") -> CanonicalDocument:
        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise ParserError(f"Failed to open DOCX stream: {e}") from e

        doc_id = document_id or str(uuid.uuid4())
        page_data = PageData(page_number=1)

        # 1. Extract Paragraphs & Headings
        current_section = ""
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""
            if "heading" in style_name:
                current_section = text
                block_type = "heading"
            else:
                block_type = "text"

            page_data.blocks.append(
                ContentBlock(
                    block_id=str(uuid.uuid4()),
                    block_type=block_type,
                    text=text,
                    section_hierarchy=[current_section] if current_section else [],
                )
            )

        # 2. Extract Tables
        for table in doc.tables:
            if not table.rows:
                continue

            headers = [cell.text.strip() for cell in table.rows[0].cells]
            rows: list[list[str]] = []
            for row in table.rows[1:]:
                row_vals = [cell.text.strip() for cell in row.cells]
                rows.append(row_vals)

            page_data.tables.append(
                TableData(
                    table_id=str(uuid.uuid4()),
                    headers=headers,
                    rows=rows,
                    page_number=1,
                )
            )

        return CanonicalDocument(
            document_id=doc_id,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            total_pages=1,
            pages=[page_data],
            metadata={"format": "DOCX", "paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)},
        )

