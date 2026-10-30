"""Unit Tests for Phase 6: Document Parsers (PDF, DOCX, PPTX, XLSX, CSV, Code)."""

import io
import docx
import fitz
import openpyxl
import pytest
from pptx import Presentation

from app.services.parser.docx_parser import DocxParser
from app.services.parser.pdf_parser import PDFParser
from app.services.parser.pptx_parser import PPTXParser
from app.services.parser.registry import default_parser_registry, parse_document_content
from app.services.parser.spreadsheet_parser import SpreadsheetParser
from app.services.parser.text_code_parser import TextCodeParser


def test_pdf_parser_with_fitz():
    # Dynamically generate a valid 2-page PDF in memory
    doc = fitz.open()
    page1 = doc.new_page(width=612, height=792)
    page1.insert_text((50, 50), "OmniRAG Executive Summary", fontsize=18)
    page1.insert_text((50, 100), "Enterprise Multimodal Knowledge Retrieval Engine.", fontsize=12)

    page2 = doc.new_page(width=612, height=792)
    page2.insert_text((50, 50), "Architecture Specifications", fontsize=14)
    page2.insert_text((50, 80), "Hybrid search with reciprocal rank fusion.", fontsize=12)

    pdf_bytes = doc.tobytes()
    doc.close()

    parser = PDFParser()
    assert parser.can_parse("application/pdf") is True

    canonical = parser.parse(pdf_bytes, "architecture.pdf")
    assert canonical.total_pages == 2
    assert len(canonical.pages) == 2
    assert "OmniRAG Executive Summary" in canonical.pages[0].blocks[0].text
    assert canonical.pages[0].blocks[0].bounding_box is not None
    assert canonical.pages[0].blocks[0].bounding_box.x0 >= 0


def test_docx_parser():
    doc = docx.Document()
    doc.add_heading("Cloud Platform Security", level=1)
    doc.add_paragraph("Zero-trust multi-tenant data segregation architecture.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Permissions"
    table.rows[1].cells[0].text = "Admin"
    table.rows[1].cells[1].text = "Full Access"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parser = DocxParser()
    canonical = parser.parse(docx_bytes, "security.docx")
    assert canonical.total_pages == 1
    assert any("Cloud Platform Security" in b.text for b in canonical.pages[0].blocks)
    assert len(canonical.pages[0].tables) == 1
    assert canonical.pages[0].tables[0].headers == ["Role", "Permissions"]
    assert canonical.pages[0].tables[0].rows == [["Admin", "Full Access"]]


def test_pptx_parser():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "OmniRAG Keynote"
    subtitle = slide.placeholders[1]
    subtitle.text = "Next Generation Multimodal Enterprise Intelligence"

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    parser = PPTXParser()
    canonical = parser.parse(pptx_bytes, "presentation.pptx")
    assert canonical.total_pages == 1
    full_text = canonical.get_full_text()
    assert "OmniRAG Keynote" in full_text


def test_spreadsheet_parser_csv_and_xlsx():
    # Test CSV
    csv_bytes = b"Product,Units,Revenue\nVectorDB,150,45000\nReranker,75,18750\n"
    parser = SpreadsheetParser()
    canonical_csv = parser.parse(csv_bytes, "metrics.csv")
    assert len(canonical_csv.pages[0].tables) == 1
    assert canonical_csv.pages[0].tables[0].headers == ["Product", "Units", "Revenue"]
    assert canonical_csv.pages[0].tables[0].rows[0] == ["VectorDB", "150", "45000"]

    # Test XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3_Performance"
    ws.append(["Metric", "Target", "Actual"])
    ws.append(["Recall@5", "0.85", "0.89"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    canonical_xlsx = parser.parse(xlsx_bytes, "perf.xlsx")
    assert len(canonical_xlsx.pages) == 1
    assert canonical_xlsx.pages[0].tables[0].headers == ["Metric", "Target", "Actual"]


def test_text_code_parser():
    md_content = """# System Guidelines
This is a standard paragraph.

def calculate_rrf(rank_dense, rank_sparse):
    return 1.0 / (60 + rank_dense) + 1.0 / (60 + rank_sparse)
""".encode("utf-8")

    parser = TextCodeParser()
    canonical = parser.parse(md_content, "algorithms.md")
    assert len(canonical.pages[0].blocks) >= 3
    assert any(b.block_type == "heading" for b in canonical.pages[0].blocks)
    assert any(b.block_type == "code" for b in canonical.pages[0].blocks)


def test_parser_registry_dispatch():
    txt_bytes = b"Simple document text"
    canonical = parse_document_content(txt_bytes, "readme.txt", "text/plain")
    assert canonical.total_pages == 1
    assert canonical.pages[0].blocks[0].text == "Simple document text"
